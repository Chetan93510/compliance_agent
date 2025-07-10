from abc import ABC, abstractmethod
import pdfplumber
from docx import Document
from PIL import Image
import pytesseract
import openpyxl
import csv

class DocumentParser(ABC):
    def _init_(self, filepath):
        self.filepath = filepath

    @abstractmethod
    def parse(self):
        pass
class PDFParser(DocumentParser):
    def parse(self):
        with pdfplumber.open(self.filepath) as pdf:
            text = ""
            for page in pdf.pages:
                text += page.extract_text()
        return text
class DocxParser(DocumentParser):
     def parse(self):
        doc = Document(self.filepath)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text
class ImageParser(DocumentParser):
        def parse(self):
         image = Image.open(self.filepath)
         text = pytesseract.image_to_string(image)
         return text
class XLSXParser(DocumentParser):
    def parse(self):
        wb = openpyxl.load_workbook(self.filepath)
        text = ""
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
        return text
class CSVParser(DocumentParser):
     def parse(self):
        with open(self.filepath, newline='', encoding='utf-8') as csvfile:
            reader = csv.reader(csvfile)
            text = ""
            for row in reader:
                text += ", ".join(row) + "\n"
        return text
def get_parser(filetype, filepath):
    if filetype == "pdf":
        return PDFParser(filepath)
    elif filetype == "docx":
        return DocxParser(filepath)
    elif filetype in ["png", "jpg", "jpeg"]:
        return ImageParser(filepath)
    elif filetype == "xlsx":
        return XLSXParser(filepath)
    elif filetype == "csv":
        return CSVParser(filepath)
    else:
        raise ValueError("Unsupported file type")

# Example usage
if __name__ == "__main__":
    filetype = "pdf"  # Change this as needed
    filepath = r"C:\Users\lucky\Documents\example.pdf"
    parser = get_parser(filetype, filepath)
    content = parser.parse()
    print(content)